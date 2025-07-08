import gradio as gr
from PIL import Image
import os
import tempfile
import uuid
import traceback
from io import BytesIO
import docx
import openpyxl

from google import genai
from google.genai import types

# --- Gemini API Configuration ---
GEMINI_API_KEY = "YOUR_API_KEY"
gemini = genai.Client(api_key=GEMINI_API_KEY)

# --- Default Config ---
DEFAULT_MAX_OUTPUT_TOKENS = 500
DEFAULT_TEMPERATURE = 0.5
DEFAULT_TOP_P = 1.0
DEFAULT_STOP_SEQUENCES = ""

GEMINI_TEXT_MODEL = "gemini-1.5-flash"
GEMINI_VISION_MODEL = "gemini-2.0-flash"
GEMINI_IMAGE_GEN_MODEL = "gemini-2.0-flash-preview-image-generation"

# --- Helpers ---
def read_document_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".docx":
        doc = docx.Document(filepath)
        return "\n".join([para.text for para in doc.paragraphs])
    elif ext == ".xlsx":
        wb = openpyxl.load_workbook(filepath)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += "\t".join([str(cell) if cell else '' for cell in row]) + "\n"
        return text
    else:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

# --- Core Functions ---
def chat_with_gemini(message, chatbox, max_tokens, temp, top_p, stop_seq, sys_inst):
    try:
        stop_sequences = [s.strip() for s in stop_seq.split(",") if s.strip()]
        response = gemini.models.generate_content(
            model=GEMINI_TEXT_MODEL,
            config=types.GenerateContentConfig(
                system_instruction=sys_inst,
                temperature=temp,
                max_output_tokens=max_tokens,
                top_p=top_p,
                stop_sequences=stop_sequences,
            ),
            contents=[types.Content(role="user", parts=[types.Part(text=message)])],
        )
        reply = response.text
        chatbox.append((message, reply))
        return "", chatbox
    except Exception as e:
        traceback.print_exc()
        chatbox.append((message, f"Lỗi: {e}"))
        return "", chatbox

def summarize_document(file_obj, chatbox, max_tokens, temp, top_p, sys_inst):
    try:
        text = read_document_text(file_obj.name)
        response = gemini.models.generate_content(
            model=GEMINI_TEXT_MODEL,
            config=types.GenerateContentConfig(
                system_instruction=sys_inst,
                temperature=temp,
                max_output_tokens=max_tokens,
                top_p=top_p,
            ),
            contents=[
                types.Part(text=text),
                types.Part(text="Vui lòng tóm tắt tài liệu trên ngắn gọn, bằng tiếng Việt."),
            ],
        )
        chatbox.append(("📄", response.text))
        return chatbox
    except Exception as e:
        traceback.print_exc()
        chatbox.append(("❌", f"Lỗi tóm tắt: {e}"))
        return chatbox

def describe_image(image, chatbox, max_tokens, temp, top_p, sys_inst):
    try:
        img_bytes = BytesIO()
        Image.fromarray(image).save(img_bytes, format="JPEG")
        response = gemini.models.generate_content(
            model=GEMINI_VISION_MODEL,
            config=types.GenerateContentConfig(
                system_instruction=sys_inst,
                temperature=temp,
                max_output_tokens=max_tokens,
                top_p=top_p,
            ),
            contents=[
                types.Part.from_bytes(img_bytes.getvalue(), mime_type="image/jpeg"),
                types.Part(text="Mô tả chi tiết hình ảnh trên bằng tiếng Việt."),
            ],
        )
        chatbox.append(("🖼️", response.text))
        return chatbox
    except Exception as e:
        traceback.print_exc()
        chatbox.append(("❌", f"Lỗi mô tả ảnh: {e}"))
        return chatbox

def generate_or_edit_image(prompt, image, chatbox):
    try:
        contents = []
        if image is not None:
            img_bytes = BytesIO()
            Image.fromarray(image).save(img_bytes, format="JPEG")
            contents.append(types.Part.from_bytes(img_bytes.getvalue(), mime_type="image/jpeg"))
        contents.append(prompt)

        response = gemini.models.generate_content(
            model=GEMINI_IMAGE_GEN_MODEL,
            config=types.GenerateContentConfig(response_modalities=["TEXT", "IMAGE"]),
            contents=contents,
        )

        gen_img_path = None
        gen_text = ""
        for part in response.candidates[0].content.parts:
            if part.text:
                gen_text = part.text
            elif part.inline_data:
                img = Image.open(BytesIO(part.inline_data.data))
                tmp_path = os.path.join(tempfile.gettempdir(), f"gen_image_{uuid.uuid4().hex}.png")
                img.save(tmp_path)
                gen_img_path = tmp_path

        chatbox.append((prompt, "Đã tạo/chỉnh sửa ảnh."))
        if gen_img_path:
            chatbox.append(("Ảnh Kết Quả", {"path": gen_img_path}))
        if gen_text:
            chatbox.append(("Mô Tả", gen_text))
        return "", None, chatbox
    except Exception as e:
        traceback.print_exc()
        chatbox.append((prompt, f"Lỗi tạo/chỉnh sửa ảnh: {e}"))
        return "", None, chatbox

# --- Gradio App ---
with gr.Blocks(title="Gemini Multimodal Chatbot") as demo:
    gr.Markdown("# 🤖 Gemini Multimodal Chatbot - AI Tổng hợp")

    with gr.Accordion("⚙️ Cài đặt Chung", open=False):
        max_tokens = gr.Slider(1, 65000, value=DEFAULT_MAX_OUTPUT_TOKENS, step=1, label="Max Tokens")
        temperature = gr.Slider(0, 2, value=DEFAULT_TEMPERATURE, step=0.01, label="Temperature")
        top_p = gr.Slider(0, 1, value=DEFAULT_TOP_P, step=0.01, label="Top P")
        stop_sequences = gr.Textbox(label="Stop Sequences (cách nhau dấu phẩy)", value=DEFAULT_STOP_SEQUENCES)

    with gr.Tabs():
        # --- Tab Chat ---
        with gr.TabItem("💬 Chat"):
            sys_inst_chat = gr.Textbox(label="Hướng dẫn Hệ thống", value="Bạn là trợ lý AI thân thiện, trả lời bằng tiếng Việt.")
            chatbox = gr.Chatbot(layout="bubble")  # 💡 Đây là điểm chỉnh để user bên phải, bot bên trái
            msg = gr.Textbox(label="Nhập câu hỏi...")
            btn_send = gr.Button("Gửi")
            btn_clear = gr.Button("🗑️ Xóa Lịch Sử")

            btn_send.click(chat_with_gemini, [msg, chatbox, max_tokens, temperature, top_p, stop_sequences, sys_inst_chat], [msg, chatbox])
            msg.submit(chat_with_gemini, [msg, chatbox, max_tokens, temperature, top_p, stop_sequences, sys_inst_chat], [msg, chatbox])
            btn_clear.click(lambda: [], outputs=[chatbox])

        # --- Tab Tóm tắt ---
        with gr.TabItem("📄 Tóm tắt Tài liệu"):
            sys_inst_doc = gr.Textbox(label="Hướng dẫn Hệ thống", value="Bạn là trợ lý AI chuyên tóm tắt tài liệu, trả lời bằng tiếng Việt.")
            chatbox_doc = gr.Chatbot()
            file_upload = gr.File(file_types=[".pdf", ".docx", ".xlsx", ".txt"], label="Tải lên Tài liệu")
            btn_clear_doc = gr.Button("🗑️ Xóa Lịch Sử")

            file_upload.upload(summarize_document, [file_upload, chatbox_doc, max_tokens, temperature, top_p, sys_inst_doc], [chatbox_doc])
            btn_clear_doc.click(lambda: [], outputs=[chatbox_doc])

        # --- Tab Mô tả Hình ảnh ---
        with gr.TabItem("🖼️ Mô tả Hình ảnh"):
            sys_inst_img = gr.Textbox(label="Hướng dẫn Hệ thống", value="Bạn là trợ lý AI mô tả ảnh, trả lời bằng tiếng Việt.")
            chatbox_img = gr.Chatbot()
            img_upload = gr.Image(type="numpy")
            btn_clear_img = gr.Button("🗑️ Xóa Lịch Sử")

            img_upload.upload(describe_image, [img_upload, chatbox_img, max_tokens, temperature, top_p, sys_inst_img], [chatbox_img])
            btn_clear_img.click(lambda: [], outputs=[chatbox_img])

        # --- Tab Tạo/Chỉnh Ảnh ---
        with gr.TabItem("🎨 Tạo/Chỉnh Ảnh AI"):
            chatbox_gen = gr.Chatbot()
            prompt = gr.Textbox(label="Prompt mô tả...")
            img_edit = gr.Image(type="numpy", label="(Tùy chọn) Tải ảnh để chỉnh sửa")
            btn_gen = gr.Button("Tạo/Chỉnh Ảnh")
            btn_clear_gen = gr.Button("🗑️ Xóa Lịch Sử")

            btn_gen.click(generate_or_edit_image, [prompt, img_edit, chatbox_gen], [prompt, img_edit, chatbox_gen])
            btn_clear_gen.click(lambda: [], outputs=[chatbox_gen])

demo.launch(share=False)